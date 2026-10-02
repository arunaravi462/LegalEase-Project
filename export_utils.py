import re
from docx import Document
from fpdf import FPDF

def sanitize_text(text: str) -> str:
    text = re.sub(r'[“”]', '"', text)
    text = re.sub(r'[‘’]', "'", text)
    return text.encode('latin-1', 'replace').decode('latin-1')

def format_docx(text: str, doc_type: str) -> str:
    doc = Document()
    doc.add_heading(doc_type.upper(), level=0)
    for line in text.split('\n'):
        if line.strip():
            doc.add_paragraph(line)
    filename = f"{doc_type.replace(' ', '_')}.docx"
    doc.save(filename)
    return filename

def format_pdf(text: str, doc_type: str) -> str:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)
    clean_text = sanitize_text(text)
    
    for line in clean_text.split('\n'):
        pdf.multi_cell(0, 8, txt=line)
        
    filename = f"{doc_type.replace(' ', '_')}.pdf"
    pdf.output(filename)
    return filename
